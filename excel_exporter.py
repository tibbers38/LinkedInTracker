from datetime import datetime
import sqlite3
from openpyxl.utils import get_column_letter

import pandas as pd
from openpyxl import Workbook, worksheet
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.styles import Border, Side
from openpyxl.styles import PatternFill
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils.dataframe import dataframe_to_rows
from openpyxl.chart import LineChart, Reference

class ExcelExporter:

    def __init__(self):
        self.conn = None
        self.data = None
        self.workbook = None
        self.games = []
        self.players = []

    def init_database(self):
        self.conn = sqlite3.connect("leaderboard.db")
        print("[+] Database connected")

    def load_data(self):
        self.data = pd.read_sql_query(
            """
            SELECT *
            FROM Leaderboard
            ORDER BY ActualDate, Game, Rank
            """,
            self.conn
        )
        self.games = sorted(self.data["Game"].unique())
        self.players = sorted(self.data["PlayerName"].unique())

        print(f"[+] Loaded {len(self.data)} records")

    def create_workbook(self):
        self.workbook = Workbook()

        default_sheet = self.workbook.active
        self.workbook.remove(default_sheet)

        print("[+] Workbook created")
    
    def create_game_sheet(self, game_name):
        worksheet = self.workbook.create_sheet(game_name)
        game_data = self.data[self.data["Game"] == game_name]
        game_data = game_data[
            ["ActualDate", "PlayerName", "LocalRank", "TimeElapsed", "HintFree", "MistakeFree"]
        ]
        game_data = game_data.sort_values(by=["ActualDate", "PlayerName"])
        worksheet.append(game_data.columns.tolist())
        for row in game_data.itertuples(index=False):
            worksheet.append(row)
        self.format_header(worksheet)
        self.auto_fit_columns(worksheet)
        self.add_borders(worksheet)
        self.freeze_header(worksheet)
        self.add_filter(worksheet)
        self.apply_conditional_formatting(worksheet)

    def create_summary_sheet(self):
        worksheet = self.workbook.create_sheet("Summary")

        worksheet["A1"] = "LinkedIn Games Tracker"
        worksheet["A3"] = "Report Generated"
        worksheet["B3"] = datetime.now()
        worksheet["A4"] = "Latest Game Data"
        last_update = self.data["ActualDate"].max()
        worksheet["B4"] = last_update
        worksheet["A6"] = "Total Days"
        total_days = self.data["ActualDate"].nunique()
        worksheet["B6"] = total_days
        worksheet["A7"] = "Players"
        total_players = len(self.data["PlayerName"].unique())
        worksheet["B7"] = total_players
        worksheet["A8"] = "Games"
        total_games = len(self.data["Game"].unique())
        worksheet["B8"] = total_games
        self.format_header(worksheet)
        self.auto_fit_columns(worksheet)
        self.write_overall_statistics(worksheet)
        self.write_latest_results(worksheet, start_row=3, start_col=4)
        self.write_overall_trend(worksheet)

    def get_overall_statistics(self):
        statistics = (
            self.data
            .groupby("PlayerName")
            .agg(
                Games=("PlayerName", "count"),
                Wins=("LocalRank", lambda x: (x == 1).sum()),
                AvgRank=("LocalRank", "mean"),
                AvgTime=("TimeElapsed", "mean"),
            )
        )
        statistics["AvgRank"] = statistics["AvgRank"].round(2)
        statistics["AvgTime"] = statistics["AvgTime"].round(1)
        performance_stats = (
            self.data
            .groupby("PlayerName")
            .agg(
                #WinRate=("LocalRank", lambda x: (x == 1).mean()),
                HintFreeRate=("HintFree", "mean"),
                MistakeFreeRate=("MistakeFree", "mean"),
            )
        )
        perfect_rate = (
            self.data
            .groupby("PlayerName")
            .apply(
                lambda group: (
                    (group["HintFree"] == 1)
                    &
                    (group["MistakeFree"] == 1)
                ).mean(),
                include_groups=False
            )
        )
        performance_stats["PerfectRate"] = perfect_rate
        statistics = statistics.join(performance_stats)
        return statistics
    
    def write_overall_statistics(self, worksheet):
        statistics = self.get_overall_statistics()
        worksheet.append([])
        worksheet.append(["Overall Statistics"])
        worksheet.append(statistics.reset_index().columns.tolist())
        for row in statistics.reset_index().itertuples(index=False):
            worksheet.append(row)

    def get_latest_results(self):
        # Lấy ngày mới nhất của từng game
        latest_dates = (
            self.data
            .groupby("Game")
            .agg(
                ActualDate=("ActualDate", "max")
            )
            .reset_index()
        )
        # Join lại với bảng gốc để lấy toàn bộ dữ liệu của ngày đó
        latest_results = latest_dates.merge(
            self.data,
            on=["Game", "ActualDate"]
        )
        # Chỉ lấy những cột cần thiết
        latest_results = latest_results[
            [
                "Game",
                "PlayerName",
                "LocalRank"
            ]
        ]
        # Pivot thành bảng
        latest_results = latest_results.pivot(
            index="PlayerName",
            columns="Game",
            values="LocalRank"
        )
        return latest_results

    def write_latest_results(self, worksheet, start_row, start_col):
        latest_results = self.get_latest_results()
        # Title
        worksheet.cell(row=start_row, column=start_col).value = "Latest Results"
        # Data
        for r, row in enumerate(
            dataframe_to_rows(latest_results.reset_index(), index=False, header=True),
            start=start_row + 1
        ):
            for c, value in enumerate(row, start=start_col):
                worksheet.cell(row=r, column=c).value = value

    def get_overall_trend(self):
        return (
            self.data
            .groupby(["ActualDate", "PlayerName"], as_index=False)
            .agg(AvgRank=("LocalRank", "mean"))
            .sort_values(["ActualDate", "PlayerName"])
        )

    def write_overall_trend(self, worksheet):
        trend = (
            self.get_overall_trend()
            .pivot(
                index="ActualDate",
                columns="PlayerName",
                values="AvgRank"
            )
            .reset_index()
        )
        start_row = 18
        start_col = 1
        # Title
        worksheet.cell(row=start_row, column=start_col).value = "Overall Trend"
        # Write table
        for r, row in enumerate(
            dataframe_to_rows(trend, index=False, header=True),
            start=start_row + 1
        ):
            for c, value in enumerate(row, start=start_col):
                worksheet.cell(row=r, column=c).value = value
        # ------------------------
        # Create chart
        # ------------------------
        chart = LineChart()
        data = Reference(
            worksheet,
            min_col=start_col + 1,
            max_col=start_col + len(trend.columns) - 1,
            min_row=start_row + 1,
            max_row=start_row + len(trend)
        )
        categories = Reference(
            worksheet,
            min_col=start_col,
            min_row=start_row + 2,
            max_row=start_row + len(trend)
        )
        chart.add_data(data, titles_from_data=True)
        chart.set_categories(categories)
        chart.title = "Average Rank Trend"
        chart.y_axis.title = "Average Rank"
        chart.x_axis.title = "Date"
        chart.y_axis.scaling.orientation = "maxMin"
        worksheet.add_chart(chart, "G18")

    def get_column_letter_by_header(self, worksheet, header_row, header_name):
        for cell in worksheet[header_row]:
            if cell.value == header_name:
                return get_column_letter(cell.column)
        return None

    def format_percentage_columns(self, worksheet, header_row, headers):
        for header in headers:
            column = self.get_column_letter_by_header(
                worksheet,
                header_row,
                header
            )
            if column is None:
                continue
            for row in range(header_row + 1, worksheet.max_row + 1):
                worksheet[f"{column}{row}"].number_format = "0%"

    # These function are for formatting the Excel sheet, such as setting header styles, auto-fitting columns, adding borders, freezing the header row, adding filters, and applying conditional formatting.

    def format_header(self, worksheet):
        header_font = Font(
            bold=True,
            color="FFFFFF"
        )
        header_fill = PatternFill(
            fill_type="solid",
            fgColor="4472C4"
        )
        header_alignment = Alignment(
            horizontal="center",
            vertical="center"
        )
        for cell in worksheet[1]:
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_alignment

    def auto_fit_columns(self, worksheet):
        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                if cell.value is not None:
                    max_length = max(
                        max_length,
                        len(str(cell.value))
                    )
            worksheet.column_dimensions[column_letter].width = max_length + 3
    
    def add_borders(self, worksheet):
        thin = Side(
            style="thin",
            color="DDDDDD"
        )
        border = Border(
            left=thin,
            right=thin,
            top=thin,
            bottom=thin
        )
        for row in worksheet.iter_rows():
            for cell in row:
                cell.border = border

    def freeze_header(self, worksheet):
        worksheet.freeze_panes = "A2"
    
    def add_filter(self, worksheet):
        worksheet.auto_filter.ref = worksheet.dimensions

    def apply_conditional_formatting(self, worksheet):
        gold = PatternFill(
            fill_type="solid",
            fgColor="FFD700"
        )
        silver = PatternFill(
            fill_type="solid",
            fgColor="D9D9D9"
        )
        bronze = PatternFill(
            fill_type="solid",
            fgColor="F4B183"
        )
        worksheet.conditional_formatting.add(
            "C2:C1000",
            CellIsRule(
                operator="equal",
                formula=["1"],
                fill=gold
            )
        )
        worksheet.conditional_formatting.add(
            "C2:C1000",
            CellIsRule(
                operator="equal",
                formula=["2"],
                fill=silver
            )
        )
        worksheet.conditional_formatting.add(
            "C2:C1000",
            CellIsRule(
                operator="equal",
                formula=["3"],
                fill=bronze
            )
        )

    def save_workbook(self):
        self.workbook.save("result.xlsx")
        print("[+] Excel exported")

    def run(self):
        self.init_database()
        self.load_data()
        self.create_workbook()
        self.create_summary_sheet()
        for game in self.games:
            self.create_game_sheet(game)
        self.save_workbook()


if __name__ == "__main__":
    exporter = ExcelExporter()
    exporter.run()